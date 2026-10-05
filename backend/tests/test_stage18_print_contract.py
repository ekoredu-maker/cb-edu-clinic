import zipfile

from openpyxl import load_workbook

from domain.settlement import build_settlement
from services.excel_service import create_execution_xlsx, create_manager_book_xlsx, create_pay_slip_xlsx
from services.hwpx_native_service import create_native_hwpx
from services.print_contract import get_print_contract
from services.print_document_model import (
    execution_report_print_model,
    manager_book_print_model,
    pay_slip_print_model,
)


def _state():
    return {
        "cfg": {
            "org": "제천교육지원청",
            "confirmer": "담당 장학사",
            "maskMode": "partial",
            "rateCoach": 40000,
            "rateClass": 30000,
            "taxPct": 3.3,
            "budget": {"total": 1000000, "coach": 600000, "cls": 300000, "travel": 100000},
        },
        "stf": [{"id": "sf001", "nm": "김지원", "ph": "010-0000-0000"}],
        "stu": [{"id": "st001", "nm": "홍길동", "sc": "의림초", "gr": 4, "cls": 2}],
        "mat": [
            {
                "id": "m1", "stfId": "sf001", "stuId": "st001", "kind": "coach", "st": "active",
                "logs": [{"date": "2026-10-05", "time": "15:00~15:50", "topic": "수학 지도", "status": "paid", "kind": "coach", "minutes": 50}],
            },
            {
                "id": "m2", "stfId": "sf001", "kind": "class", "st": "active",
                "classInfo": {"sc": "의림초", "scType": "초", "gr": 3, "cls": 1},
                "logs": [{"date": "2026-10-06", "time": "09:00~09:40", "topic": "수업협력", "status": "verified", "kind": "class", "minutes": 40}],
            },
        ],
        "trn": [],
    }


def _col_counts(path):
    with zipfile.ZipFile(path, "r") as zin:
        section = zin.read("Contents/section0.xml").decode("utf-8")
    import re
    return [int(x) for x in re.findall(r'<hp:tbl\b[^>]*\bcolCnt="(\d+)"', section)]


def test_stage18_print_contracts_match_html_print_tables(tmp_path):
    state = _state()
    settlement = build_settlement({"state": state, "ym": "2026-10"})
    staff = state["stf"][0]
    staff_map = {staff["id"]: staff}

    pay = pay_slip_print_model(settlement, staff, "sf001", "2026-10", state["cfg"]["org"])
    execution = execution_report_print_model(settlement, staff_map, "2026-10", state["cfg"]["org"])
    coach = manager_book_print_model(state, "sf001", "2026-10", kind="coach")
    cls = manager_book_print_model(state, "sf001", "2026-10", kind="class")

    assert pay["columns"] == get_print_contract("pay_slip")["columns"]
    assert len(pay["columns"]) == 6
    assert all(len(row) == 6 for row in pay["rows"])

    assert execution["columns"] == get_print_contract("execution_report")["columns"]
    assert len(execution["columns"]) == 9
    assert all(len(row) == 9 for row in execution["rows"])
    assert len(execution["summaryRow"]) == 9

    assert len(coach["columns"]) == 12
    assert coach["columns"][5] == coach["columns"][11] == "학생"
    assert dict(coach["meta"])["구분"] == "학습코칭"
    assert "— 학습코칭" in coach["title"]
    assert coach["rows"][4][5] == "홍OO"

    assert len(cls["columns"]) == 12
    assert cls["columns"][5] == cls["columns"][11] == "학급"
    assert dict(cls["meta"])["구분"] == "수업협력"
    assert "— 수업협력" in cls["title"]
    assert cls["rows"][5][5] == "초 3-1반"

    pay_path = create_native_hwpx(tmp_path / "pay.hwpx", pay)
    exec_path = create_native_hwpx(tmp_path / "exec.hwpx", execution)
    mgr_path = create_native_hwpx(tmp_path / "manager.hwpx", cls)
    assert _col_counts(pay_path)[0] == 6
    assert _col_counts(exec_path)[0] == 9
    assert _col_counts(mgr_path)[0] == 12


def test_stage18_excel_keeps_detailed_export_schema(tmp_path):
    state = _state()
    settlement = build_settlement({"state": state, "ym": "2026-10"})
    staff = state["stf"][0]
    staff_map = {staff["id"]: staff}

    pay_path = create_pay_slip_xlsx(tmp_path / "pay.xlsx", settlement, staff, "sf001", "2026-10", state["cfg"]["org"])
    exec_path = create_execution_xlsx(tmp_path / "exec.xlsx", settlement, staff_map, "2026-10", state["cfg"]["org"])
    mgr_path = create_manager_book_xlsx(tmp_path / "mgr.xlsx", state, "sf001", "2026-10")

    pay_ws = load_workbook(pay_path, data_only=True)["지급명세서"]
    exec_ws = load_workbook(exec_path, data_only=True)["집행내역서"]
    mgr_ws = load_workbook(mgr_path, data_only=True)["관리부"]

    assert len([pay_ws.cell(4, c).value for c in range(1, 9)]) == 8
    assert pay_ws.cell(4, 2).value == "구분"
    assert len([exec_ws.cell(4, c).value for c in range(1, 11)]) == 10
    assert exec_ws.cell(4, 10).value == "비고"
    assert len([mgr_ws.cell(5, c).value for c in range(1, 14)]) == 13
    assert mgr_ws.cell(5, 7).value in (None, "")
