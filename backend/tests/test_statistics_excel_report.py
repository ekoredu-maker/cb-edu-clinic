from openpyxl import load_workbook

from services.statistics_excel_service import create_statistics_report_xlsx
from tests.load_fixture import build_large_state


def test_statistics_report_is_administrative_workbook(tmp_path):
    state = build_large_state(ym="2026-10")
    state["cfg"].update({
        "org": "제천교육지원청",
        "base": "제천거점",
        "admin": "테스트 담당자",
        "regions": ["제천", "단양"],
    })

    path = create_statistics_report_xlsx(
        tmp_path / "statistics_report.xlsx",
        state,
        as_of="2026-10-05",
        report_type="base",
    )

    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == ["종합보고서", "방과후학습코칭", "치료기관연계", "난독경계선"]

    summary = wb["종합보고서"]
    assert "제천교육지원청 제천거점 지원 실적" in str(summary["A1"].value)
    assert summary["A3"].value and "2026.10.05" in str(summary["A3"].value)
    assert summary["A6"].value == "등록 학생"
    assert summary["A7"].value == "1,500명"
    assert summary["C6"].value == "활동 지원단"
    assert summary.page_setup.orientation == "landscape"
    assert summary.page_setup.fitToWidth == 1
    assert "A1:J2" in {str(rng) for rng in summary.merged_cells.ranges}
    assert summary.print_area

    therapy = wb["치료기관연계"]
    headers = [therapy.cell(row=5, column=i).value for i in range(1, 9)]
    assert headers == ["번호", "학교명", "학년", "성별", "지원내용", "연계기관", "지원기간", "비고"]
    assert "학생명" not in headers
    assert therapy.page_setup.orientation == "landscape"

    after_school = wb["방과후학습코칭"]
    assert "방과후학습코칭 지원현황" in str(after_school["A1"].value)
    assert after_school.page_setup.fitToWidth == 1

    dyslex = wb["난독경계선"]
    assert "난독증 및 경계선지능 지원 현황" in str(dyslex["A1"].value)
