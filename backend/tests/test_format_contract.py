from openpyxl import load_workbook

from domain.settlement import build_settlement
from services.excel_service import create_execution_xlsx, create_manager_book_xlsx, create_pay_slip_xlsx
from services.format_contract import get_format_contract, list_format_contracts
from tests.load_fixture import build_large_state


def test_embedded_html_format_contracts_are_registered():
    contracts = list_format_contracts()
    assert set(contracts) == {"pay_slip", "execution_report", "manager_book"}

    pay = get_format_contract("pay_slip")
    assert pay["source"]["file"] == "legacy-js/10-ext-v99.js"
    assert pay["columns"] == ["날짜", "구분", "학생/제목", "학교", "시간/회차", "내용", "단가", "지급액"]
    assert pay["columnWidths"] == [12, 10, 16, 16, 14, 22, 12, 12]

    execution = get_format_contract("execution_report")
    assert execution["columns"][-1] == "비고"
    assert execution["sheet"] == "집행내역서"

    manager = get_format_contract("manager_book")
    assert manager["layout"]["rows"] == 16
    assert manager["layout"]["leftDays"] == [1, 16]
    assert manager["layout"]["rightDays"] == [17, 31]
    assert manager["verifiedStatuses"] == ["verified", "paid"]


def test_python_excel_outputs_follow_html_contract(tmp_path):
    state = build_large_state(ym="2026-10")
    settlement = build_settlement({"state": state, "ym": "2026-10"})
    staff_by_id = {str(x["id"]): x for x in state["stf"]}
    staff_id = "sf001"

    pay_path = tmp_path / "pay.xlsx"
    create_pay_slip_xlsx(pay_path, settlement, staff_by_id[staff_id], staff_id, "2026-10", state["cfg"]["org"])
    pay_ws = load_workbook(pay_path, data_only=True)["지급명세서"]
    assert pay_ws["A1"].value == "활동비 지급 명세서 (2026-10)"
    assert [pay_ws.cell(4, c).value for c in range(1, 9)] == get_format_contract("pay_slip")["columns"]
    assert "수령인: 지원단001" in pay_ws["A2"].value

    exec_path = tmp_path / "execution.xlsx"
    create_execution_xlsx(exec_path, settlement, staff_by_id, "2026-10", state["cfg"]["org"])
    exec_ws = load_workbook(exec_path, data_only=True)["집행내역서"]
    assert exec_ws["A1"].value == "월별 활동비 집행내역서 (2026-10)"
    assert [exec_ws.cell(4, c).value for c in range(1, 11)] == get_format_contract("execution_report")["columns"]
    assert exec_ws["J5"].value in (None, "")

    mgr_path = tmp_path / "manager.xlsx"
    create_manager_book_xlsx(mgr_path, state, staff_id, "2026-10")
    mgr_ws = load_workbook(mgr_path, data_only=True)["관리부"]
    assert mgr_ws["A1"].value == "학습지원단 관리부 (2026년 10월)"
    assert "담당학생:" in mgr_ws["A3"].value
    assert [mgr_ws.cell(5, c).value or "" for c in range(1, 14)] == get_format_contract("manager_book")["columns"]
    assert mgr_ws.max_row >= 22
    summary_values = [mgr_ws.cell(mgr_ws.max_row, c).value for c in range(1, 14)]
    assert any("총 실시 회기: 120회" in str(v) for v in summary_values if v is not None)
