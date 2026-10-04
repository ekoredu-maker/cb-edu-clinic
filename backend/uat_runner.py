from __future__ import annotations

import argparse
import hashlib
import json
import tempfile
from copy import deepcopy
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Callable

from openpyxl import load_workbook

import database.db as db_module
import database.state_store as store_module
from domain.settlement import build_settlement
from domain.statistics import build_statistics
from domain.verification import verify_records
from services.excel_service import create_execution_xlsx, create_pay_slip_xlsx
from tests.load_fixture import build_large_state


YM = "2026-10"


def _canonical(value: Any) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), default=str)


def _sha256(value: Any) -> str:
    return hashlib.sha256(_canonical(value).encode("utf-8")).hexdigest()


def _use_db(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    db_module.DATA_DIR = path.parent
    db_module.DB_PATH = path
    store_module.connect = db_module.connect
    db_module.init_db()


def _calc_snapshot(state: dict[str, Any]) -> dict[str, Any]:
    stats = build_statistics({"state": state})
    verification = verify_records({"state": state, "ym": YM})
    settlement = build_settlement({"state": state, "ym": YM})
    return {
        "statistics": {
            "students": stats["counts"]["students"],
            "staff": stats["counts"]["staff"],
            "matchings": stats["counts"]["matchings"],
            "actualCoachStudents": stats["counts"]["actualCoachStudents"],
            "actualClassMatchings": stats["counts"]["actualClassMatchings"],
        },
        "verification": {
            "errorCount": verification["errorCount"],
            "monthlyCount": len(verification["monthly"]),
        },
        "settlement": {
            "coach": settlement["executed"]["coach"],
            "cls": settlement["executed"]["cls"],
            "travel": settlement["executed"]["travel"],
            "total": settlement["executed"]["total"],
        },
    }


def _assert_equal(actual: Any, expected: Any, label: str) -> None:
    if actual != expected:
        raise AssertionError(f"{label}: expected={expected!r}, actual={actual!r}")


def _find_row(ws, first_cell_value: str) -> int:
    for row in range(1, ws.max_row + 1):
        if ws.cell(row, 1).value == first_cell_value:
            return row
    raise AssertionError(f"Excel에서 '{first_cell_value}' 행을 찾지 못했습니다.")


def _find_row_any_cell(ws, value: str) -> int:
    for row in range(1, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            if ws.cell(row, col).value == value:
                return row
    raise AssertionError(f"Excel에서 '{value}' 행을 찾지 못했습니다.")


def _write_reports(output_dir: Path, report: dict[str, Any]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    report["summary"] = {
        "ok": bool(report.get("scenarios")) and all(x["status"] == "PASS" for x in report["scenarios"]),
        "passed": sum(1 for x in report.get("scenarios", []) if x["status"] == "PASS"),
        "total": len(report.get("scenarios", [])),
    }
    (output_dir / "uat_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2, default=str), encoding="utf-8"
    )

    lines = [
        "# V13 운영규모 UAT 결과",
        "",
        f"- 생성시각(UTC): {report['generatedAt']}",
        f"- 결과: {'PASS' if report['summary']['ok'] else 'FAIL'}",
        f"- 시나리오: {report['summary']['passed']}/{report['summary']['total']} PASS",
        "",
        "## 시험 규모",
        "",
        f"- 지원단: {report['scale']['staff']}명",
        f"- 학생: {report['scale']['students']}명",
        f"- 매칭: {report['scale']['matchings']}건",
        f"- 활동로그: {report['scale']['logs']}건",
        f"- 연수: {report['scale']['trainings']}건",
        "",
        "## 시나리오",
        "",
    ]
    for idx, scenario in enumerate(report["scenarios"], 1):
        lines.append(f"{idx}. **{scenario['status']}** — {scenario['name']}")
        detail = scenario.get("detail")
        if detail:
            lines.append(f"   - `{json.dumps(detail, ensure_ascii=False, default=str)}`")
    if report.get("backupSha256"):
        lines += ["", "## 백업 무결성", "", f"- SHA-256: `{report['backupSha256']}`"]
    (output_dir / "uat_report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def run_uat(output_dir: Path) -> dict[str, Any]:
    state = build_large_state(ym=YM)
    report: dict[str, Any] = {
        "generatedAt": datetime.now(timezone.utc).isoformat(),
        "scale": {
            "staff": len(state["stf"]),
            "students": len(state["stu"]),
            "matchings": len(state["mat"]),
            "logs": sum(len(m.get("logs") or []) for m in state["mat"]),
            "trainings": len(state["trn"]),
        },
        "scenarios": [],
    }
    ctx: dict[str, Any] = {"original": state}

    def step(name: str, fn: Callable[[], Any]) -> None:
        try:
            detail = fn()
            report["scenarios"].append({"name": name, "status": "PASS", "detail": detail})
            _write_reports(output_dir, report)
        except Exception as exc:
            report["scenarios"].append({
                "name": name,
                "status": "FAIL",
                "detail": {"error": f"{type(exc).__name__}: {exc}"},
            })
            _write_reports(output_dir, report)
            raise

    with tempfile.TemporaryDirectory(prefix="cb_clinic_uat_") as tmp:
        tmp_root = Path(tmp)
        primary_db = tmp_root / "primary" / "clinic_v13.db"
        restore_db = tmp_root / "restore" / "clinic_v13.db"
        backup_path = tmp_root / "uat_backup.json"
        output_xlsx = tmp_root / "xlsx"

        def baseline_import() -> dict[str, Any]:
            _use_db(primary_db)
            baseline = _calc_snapshot(state)
            _assert_equal(baseline["statistics"]["students"], 1500, "학생 수")
            _assert_equal(baseline["statistics"]["staff"], 50, "지원단 수")
            _assert_equal(baseline["statistics"]["matchings"], 1500, "매칭 수")
            _assert_equal(baseline["verification"]["errorCount"], 0, "검증 오류")
            _assert_equal(baseline["verification"]["monthlyCount"], 1500, "월 검증 대상")
            _assert_equal(baseline["settlement"], {
                "coach": 192_000_000,
                "cls": 36_000_000,
                "travel": 18_000_000,
                "total": 246_000_000,
            }, "초기 정산")
            imported = store_module.import_state(state, source="uat-initial", replace=True)
            _assert_equal(imported["recordCount"], 3074, "SQLite 초기 이관 레코드")
            if not store_module.compare_state(state)["ok"]:
                raise AssertionError("초기 Browser↔SQLite parity 실패")
            ctx["baselineCalc"] = baseline
            return {"recordCount": imported["recordCount"], "settlement": baseline["settlement"]}

        step("운영규모 초기 이관 및 기준 계산", baseline_import)

        def mutate_and_dual_write() -> dict[str, Any]:
            working = deepcopy(state)
            for i in range(100):
                working["stu"][i]["memo"] = f"UAT학생수정-{i}"
                store_module.upsert_record("stu", working["stu"][i], source="uat-edit")
            for i in range(50):
                working["mat"][i]["memo"] = f"UAT매칭수정-{i}"
                store_module.upsert_record("mat", working["mat"][i], source="uat-edit")

            new_student = {
                "id": "stUAT1501",
                "nm": "UAT신규학생",
                "sc": "가상학교61",
                "scType": "초",
                "gr": 3,
                "region": "제천",
                "supportTypes": ["방과후학습코칭"],
                "st": "active",
            }
            working["stu"].append(new_student)
            store_module.upsert_record("stu", new_student, source="uat-add")

            deleted_ids = [row["id"] for row in working["trn"][-4:]]
            for row_id in deleted_ids:
                store_module.delete_record("trn", row_id, source="uat-delete")
            working["trn"] = working["trn"][:-4]
            working["cfg"]["confirmer"] = "UAT 변경확인자"
            store_module.save_singleton("cfg", working["cfg"], source="uat-settings")

            parity = store_module.compare_state(working)
            if not parity["ok"]:
                raise AssertionError(f"수정 후 parity 실패: {parity}")

            calc = _calc_snapshot(working)
            _assert_equal(calc["statistics"]["students"], 1501, "수정 후 학생 수")
            _assert_equal(calc["statistics"]["matchings"], 1500, "수정 후 매칭 수")
            _assert_equal(calc["verification"]["monthlyCount"], 1500, "수정 후 월 검증 대상")
            _assert_equal(calc["settlement"], {
                "coach": 192_000_000,
                "cls": 36_000_000,
                "travel": 15_000_000,
                "total": 243_000_000,
            }, "수정 후 정산")
            ctx["working"] = working
            ctx["workingCalc"] = calc
            return {
                "studentUpdated": 100,
                "matchingUpdated": 50,
                "studentAdded": 1,
                "trainingDeleted": deleted_ids,
                "settlement": calc["settlement"],
            }

        step("수정·추가·삭제·설정변경 및 SQLite 이중쓰기", mutate_and_dual_write)

        def restart_reload() -> dict[str, Any]:
            db_module.init_db()
            reloaded = store_module.export_state()
            _assert_equal(_sha256(reloaded), _sha256(ctx["working"]), "재기동 후 상태 해시")
            if not store_module.compare_state(ctx["working"])["ok"]:
                raise AssertionError("재기동 후 Browser↔SQLite parity 실패")
            calc = _calc_snapshot(reloaded)
            _assert_equal(calc, ctx["workingCalc"], "재기동 후 계산 결과")
            with db_module.connect() as conn:
                audit_count = int(conn.execute("SELECT COUNT(*) FROM audit_log").fetchone()[0])
            if audit_count < 150:
                raise AssertionError(f"감사로그가 예상보다 적습니다: {audit_count}")
            ctx["reloaded"] = reloaded
            return {"stateSha256": _sha256(reloaded), "auditRows": audit_count}

        step("종료·재실행에 준하는 SQLite 재로드", restart_reload)

        def backup_json() -> dict[str, Any]:
            exported = store_module.export_state()
            backup_text = json.dumps(exported, ensure_ascii=False, indent=2, default=str)
            backup_path.write_text(backup_text, encoding="utf-8")
            parsed = json.loads(backup_path.read_text(encoding="utf-8"))
            _assert_equal(_sha256(parsed), _sha256(ctx["working"]), "백업 JSON 상태 해시")
            raw_sha = hashlib.sha256(backup_path.read_bytes()).hexdigest()
            report["backupSha256"] = raw_sha
            ctx["backupState"] = parsed
            return {"bytes": backup_path.stat().st_size, "sha256": raw_sha}

        step("JSON 백업 생성 및 무결성 검증", backup_json)

        def detect_and_recover_drift() -> dict[str, Any]:
            corrupted = deepcopy(ctx["working"]["stu"][777])
            corrupted["nm"] = "의도적_DB_손상"
            store_module.upsert_record("stu", corrupted, source="uat-forced-drift")
            diff = store_module.compare_state(ctx["working"])
            if diff["ok"]:
                raise AssertionError("의도적 DB 손상을 감지하지 못했습니다.")
            stu_check = next(x for x in diff["checks"] if x["key"] == "stu")
            expected_id = ctx["working"]["stu"][777]["id"]
            _assert_equal(stu_check["changed"], [expected_id], "손상 레코드 식별")

            store_module.import_state(ctx["working"], source="uat-browser-recovery", replace=True)
            recovered = store_module.compare_state(ctx["working"])
            if not recovered["ok"]:
                raise AssertionError("Browser 기준 SQLite 복구 실패")
            _assert_equal(_calc_snapshot(store_module.export_state()), ctx["workingCalc"], "복구 후 계산")
            return {"detectedChangedId": expected_id, "recovery": "PASS"}

        step("의도적 DB 손상 감지 및 Browser 기준 복구", detect_and_recover_drift)

        def restore_into_fresh_db() -> dict[str, Any]:
            _use_db(restore_db)
            imported = store_module.import_state(ctx["backupState"], source="uat-backup-restore", replace=True)
            if not store_module.compare_state(ctx["backupState"])["ok"]:
                raise AssertionError("새 DB에 백업 복원 후 parity 실패")
            restored = store_module.export_state()
            _assert_equal(_sha256(restored), _sha256(ctx["backupState"]), "새 DB 복원 상태 해시")
            restored_calc = _calc_snapshot(restored)
            _assert_equal(restored_calc, ctx["workingCalc"], "백업 복원 후 계산 결과")
            ctx["restored"] = restored
            ctx["restoredCalc"] = restored_calc
            return {"recordCount": imported["recordCount"], "stateSha256": _sha256(restored)}

        step("새 SQLite DB에 백업 복원 및 계산 재검증", restore_into_fresh_db)

        def verify_excel_outputs() -> dict[str, Any]:
            restored = ctx["restored"]
            settlement = build_settlement({"state": restored, "ym": YM})
            staff_by_id = {str(x["id"]): x for x in restored["stf"]}
            execution_path = output_xlsx / "월별집행내역_UAT.xlsx"
            create_execution_xlsx(
                execution_path,
                settlement,
                staff_by_id,
                YM,
                restored["cfg"].get("org") or "",
            )
            wb = load_workbook(execution_path, data_only=True)
            ws = wb["집행내역서"]
            total_row = _find_row_any_cell(ws, "합계")
            ex = settlement["executed"]
            _assert_equal(ws.cell(total_row, 4).value, ex["coach"], "집행내역 코칭 합계")
            _assert_equal(ws.cell(total_row, 6).value, ex["cls"], "집행내역 협력 합계")
            _assert_equal(ws.cell(total_row, 8).value, ex["travel"], "집행내역 출장 합계")
            _assert_equal(ws.cell(total_row, 9).value, ex["total"], "집행내역 총합계")

            staff_id = restored["stf"][0]["id"]
            staff = staff_by_id[staff_id]
            pay_path = output_xlsx / "지급명세서_UAT.xlsx"
            create_pay_slip_xlsx(
                pay_path,
                settlement,
                staff,
                staff_id,
                YM,
                restored["cfg"].get("org") or "",
            )
            pay_wb = load_workbook(pay_path, data_only=True)
            pay_ws = pay_wb["지급명세서"]
            pay_row = _find_row(pay_ws, "단가 합계 / 지급액(총액)")
            summary = settlement["summaryByStaff"][staff_id]
            _assert_equal(pay_ws.cell(pay_row, 2).value, summary["gross"], "지급명세서 세전")
            _assert_equal(pay_ws.cell(pay_row, 4).value, -summary["tax"], "지급명세서 공제")
            _assert_equal(pay_ws.cell(pay_row, 6).value, summary["net"], "지급명세서 실지급")
            return {
                "executionTotal": ex["total"],
                "paySlipStaff": staff["nm"],
                "paySlipGross": summary["gross"],
                "paySlipNet": summary["net"],
            }

        step("복원 상태의 정산·Excel 출력 합계 대조", verify_excel_outputs)

    _write_reports(output_dir, report)
    return report


def main() -> int:
    parser = argparse.ArgumentParser(description="학습클리닉 V13 운영규모 UAT")
    parser.add_argument("--output-dir", default="../.build/uat", help="UAT 보고서 출력 폴더")
    args = parser.parse_args()
    report = run_uat(Path(args.output_dir).resolve())
    summary = report["summary"]
    print(f"V13 UAT: {summary['passed']}/{summary['total']} PASS")
    print(f"Report: {Path(args.output_dir).resolve() / 'uat_report.md'}")
    return 0 if summary["ok"] else 1


if __name__ == "__main__":
    raise SystemExit(main())
