from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from domain.statistics import build_statistics


AREA_SPECS = [
    ("HANGUL", "한글미해득", "한글미해득"),
    ("BASIC", "기초학습지원", "기초학습지원"),
    ("DYSLEX", "난독증", "난독증"),
    ("READING", "읽기곤란", "읽기곤란"),
    ("ADHD", "ADHD", "ADHD"),
    ("BORDER", "경계선지능", "경계선지능"),
    ("EMOTION", "심리정서", "심리정서"),
    ("LANG", "언어발달지연", "언어"),
    ("ETC", "기타", "기타"),
]
AFTER_SCHOOL_AREAS = [x for x in AREA_SPECS if x[0] != "LANG"]
THERAPY_AREA_IDS = ["DYSLEX", "LANG", "BORDER", "ADHD", "EMOTION", "ETC"]
AREA_BY_ID = {key: {"label": label, "alt": alt} for key, label, alt in AREA_SPECS}

NAVY = "1F4E78"
BLUE = "D9EAF7"
LIGHT_BLUE = "EEF5FA"
GRAY = "F3F4F6"
DARK = "1F2937"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
YELLOW = "FFF2CC"
BORDER_COLOR = "8A99A8"
THIN = Side(style="thin", color=BORDER_COLOR)
MEDIUM = Side(style="medium", color=NAVY)


def _students(state: dict[str, Any]) -> list[dict[str, Any]]:
    return state.get("stu") or state.get("students") or []


def _regions(state: dict[str, Any]) -> list[str]:
    cfg = state.get("cfg") or {}
    regions = list(cfg.get("regions") or [])
    return regions or ["지역1", "지역2"]


def _today_text(value: str | None) -> str:
    raw = str(value or "").strip()
    if raw:
        return raw.replace("-", ".")
    return date.today().isoformat().replace("-", ".")


def _school_grade_keys() -> list[tuple[str, int]]:
    return [("초", i) for i in range(1, 7)] + [("중", i) for i in range(1, 4)]


def _after_school_pivot(state: dict[str, Any]) -> dict[str, dict[str, int]]:
    ids = [x[0] for x in AFTER_SCHOOL_AREAS]
    result = {f"{lv}{gr}": {area: 0 for area in ids} for lv, gr in _school_grade_keys()}
    for student in _students(state):
        if "방과후학습코칭" not in (student.get("supportTypes") or []):
            continue
        key = f"{student.get('scType', '')}{student.get('gr', '')}"
        row = result.get(key)
        if not row:
            continue
        for area in student.get("areas") or []:
            if area in row:
                row[area] += 1
    return result


def _therapy_data(state: dict[str, Any]) -> tuple[list[dict[str, Any]], dict[str, dict[str, int]]]:
    students = [s for s in _students(state) if "치료기관연계" in (s.get("supportTypes") or [])]
    pivot = {f"{lv}{gr}": {area: 0 for area in THERAPY_AREA_IDS} for lv, gr in _school_grade_keys()}
    for student in students:
        key = f"{student.get('scType', '')}{student.get('gr', '')}"
        row = pivot.get(key)
        if not row:
            continue
        for area in student.get("areas") or []:
            if area in row:
                row[area] += 1
    return students, pivot


def _dyslex_border(state: dict[str, Any]) -> dict[str, dict[str, Any]]:
    def empty() -> dict[str, Any]:
        return {
            "applied": 0, "tested": 0,
            "pos_therapy": 0, "pos_coach": 0, "pos_class": 0,
            "neg_therapy": 0, "neg_coach": 0, "neg_class": 0,
            "unsup": 0, "unsupReasons": [],
        }

    result = {"dyslex": empty(), "border": empty()}
    for student in _students(state):
        areas = student.get("areas") or []
        types = student.get("supportTypes") or []
        diag = student.get("diagTest") or {}
        unsupported = student.get("unsupported") or {}
        for key, area_id, positive_key in (
            ("dyslex", "DYSLEX", "dyslexia"),
            ("border", "BORDER", "borderline"),
        ):
            if area_id not in areas:
                continue
            row = result[key]
            row["applied"] += 1
            if diag.get("done"):
                row["tested"] += 1
            prefix = "pos" if diag.get(positive_key) else "neg"
            if "치료기관연계" in types:
                row[f"{prefix}_therapy"] += 1
            if "방과후학습코칭" in types:
                row[f"{prefix}_coach"] += 1
            if "수업협력코칭" in types:
                row[f"{prefix}_class"] += 1
            if unsupported.get("is"):
                row["unsup"] += 1
                reason = str(unsupported.get("reason") or "").strip()
                if reason:
                    row["unsupReasons"].append(reason)
    return result


def _set_border(cell, *, medium_bottom: bool = False) -> None:
    cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=MEDIUM if medium_bottom else THIN)


def _merge_title(ws, title: str, subtitle: str, max_col: int) -> None:
    end = get_column_letter(max_col)
    ws.merge_cells(f"A1:{end}2")
    cell = ws["A1"]
    cell.value = title
    cell.font = Font(name="맑은 고딕", size=18, bold=True, color=DARK)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 10
    ws.merge_cells(f"A3:{end}3")
    ws["A3"] = subtitle
    ws["A3"].font = Font(name="맑은 고딕", size=9, color="4B5563")
    ws["A3"].alignment = Alignment(horizontal="right")


def _section(ws, row: int, text: str, max_col: int) -> int:
    end = get_column_letter(max_col)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=max_col)
    c = ws.cell(row=row, column=1, value=text)
    c.fill = PatternFill("solid", fgColor=NAVY)
    c.font = Font(name="맑은 고딕", size=10, bold=True, color=WHITE)
    c.alignment = Alignment(vertical="center")
    ws.row_dimensions[row].height = 22
    return row + 1


def _style_header_row(ws, row: int, start_col: int, end_col: int, fill: str = BLUE) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="맑은 고딕", size=9, bold=True, color=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _set_border(cell)


def _style_data_range(ws, start_row: int, end_row: int, start_col: int, end_col: int) -> None:
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.font = Font(name="맑은 고딕", size=9, color=DARK)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            _set_border(cell)


def _style_total_row(ws, row: int, start_col: int, end_col: int) -> None:
    for col in range(start_col, end_col + 1):
        cell = ws.cell(row=row, column=col)
        cell.fill = PatternFill("solid", fgColor=GREEN)
        cell.font = Font(name="맑은 고딕", size=9, bold=True, color=DARK)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        _set_border(cell, medium_bottom=True)


def _setup_page(ws, *, landscape: bool = True, repeat_rows: str | None = None) -> None:
    ws.sheet_view.showGridLines = False
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE if landscape else ws.ORIENTATION_PORTRAIT
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25
    ws.page_margins.top = 0.45
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15
    ws.oddFooter.center.text = "&P / &N"
    if repeat_rows:
        ws.print_title_rows = repeat_rows


def _metric_block(ws, counts: dict[str, Any], row: int) -> int:
    metrics = [
        ("등록 학생", counts.get("students", 0), "명"),
        ("활동 지원단", counts.get("staff", 0), "명"),
        ("전체 매칭", counts.get("matchings", 0), "건"),
        ("실제 학습코칭", counts.get("actualCoachStudents", 0), "명"),
        ("실제 수업협력", counts.get("actualClassMatchings", 0), "학급"),
    ]
    col = 1
    for label, value, unit in metrics:
        ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
        ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
        h = ws.cell(row=row, column=col, value=label)
        v = ws.cell(row=row + 1, column=col, value=f"{int(value):,}{unit}")
        for c in (h, v):
            c.alignment = Alignment(horizontal="center", vertical="center")
            _set_border(c)
        h.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
        h.font = Font(name="맑은 고딕", size=8, bold=True, color=DARK)
        v.font = Font(name="맑은 고딕", size=13, bold=True, color=NAVY)
        col += 2
    ws.row_dimensions[row].height = 20
    ws.row_dimensions[row + 1].height = 27
    return row + 3


def _write_grade_table(ws, pivot: dict[str, dict[str, int]], row: int) -> int:
    headers = ["학교급", "학년", "방과후 신청", "수업협력 신청(학급)", "수업협력 신청(학생)", "심리진단", "방과후 실적", "수업협력 실적(학급)", "수업협력 실적(학생)", "치료기관연계"]
    for col, value in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=value)
    _style_header_row(ws, row, 1, 10)
    start = row + 1
    totals = {k: 0 for k in ("apply_coach", "apply_class_cnt", "apply_class_stu", "diag", "coach", "class_cnt", "class_stu", "therapy")}
    for level, max_grade in (("초", 6), ("중", 3)):
        subtotal = {k: 0 for k in totals}
        for grade in range(1, max_grade + 1):
            data = pivot.get(f"{level}{grade}") or {}
            vals = [data.get(k, 0) for k in totals]
            ws.append([level if grade == 1 else "", grade, *vals])
            for key in totals:
                subtotal[key] += int(data.get(key, 0) or 0)
                totals[key] += int(data.get(key, 0) or 0)
        ws.append([f"{level} 소계", "", *[subtotal[k] for k in totals]])
        _style_total_row(ws, ws.max_row, 1, 10)
    ws.append(["합계", "", *[totals[k] for k in totals]])
    _style_total_row(ws, ws.max_row, 1, 10)
    _style_data_range(ws, start, ws.max_row - 1, 1, 10)
    return ws.max_row + 2


def _write_region_table(ws, pivot: dict[str, dict[str, int]], regions: list[str], row: int) -> int:
    headers = ["지역", "방과후 신청", "수업협력 신청(학급)", "수업협력 신청(학생)", "심리진단", "방과후 실적", "수업협력 실적(학급)", "수업협력 실적(학생)", "치료기관연계"]
    for col, value in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=value)
    _style_header_row(ws, row, 1, 9)
    keys = ("apply_coach", "apply_class_cnt", "apply_class_stu", "diag", "coach", "class_cnt", "class_stu", "therapy")
    totals = {k: 0 for k in keys}
    start = row + 1
    for region in regions:
        data = pivot.get(region) or {}
        values = [int(data.get(k, 0) or 0) for k in keys]
        ws.append([region, *values])
        for k in keys:
            totals[k] += int(data.get(k, 0) or 0)
    ws.append(["합계", *[totals[k] for k in keys]])
    _style_data_range(ws, start, ws.max_row - 1, 1, 9)
    _style_total_row(ws, ws.max_row, 1, 9)
    return ws.max_row + 2


def _build_summary_sheet(wb: Workbook, state: dict[str, Any], statistics: dict[str, Any], as_of: str, report_type: str) -> None:
    ws = wb.active
    ws.title = "종합보고서"
    cfg = state.get("cfg") or {}
    org = str(cfg.get("org") or "충북학습종합클리닉센터")
    base = str(cfg.get("base") or "거점센터")
    admin = str(cfg.get("admin") or "")
    date_text = _today_text(as_of)
    title = f"{org} 지역거점 지원 실적" if report_type == "quarter" else f"{org} {base} 지원 실적"
    _merge_title(ws, title, f"기준일: {date_text}  |  작성: {admin or '-'}", 10)

    row = 5
    row = _section(ws, row, "Ⅰ. 핵심 현황", 10)
    row = _metric_block(ws, statistics.get("counts") or {}, row)

    row = _section(ws, row, "Ⅱ. 학교급별 지원 실적", 10)
    row = _write_grade_table(ws, statistics.get("pivotByGrade") or {}, row)

    row = _section(ws, row, "Ⅲ. 지역별 지원 현황", 10)
    row = _write_region_table(ws, statistics.get("pivotByRegion") or {}, _regions(state), row)

    row = _section(ws, row, "Ⅳ. 작성 기준", 10)
    notes = [
        "• 신청 현황은 학생 등록 및 수업협력 활성 매칭을 기준으로 집계합니다.",
        "• 실제 실적은 실시·검증·지급 상태의 활동 로그를 기준으로 집계합니다.",
        "• 치료기관연계 세부 명단은 개인정보 최소화를 위해 성명 없이 별도 시트에 제공합니다.",
        "• 이 파일은 보고·회의·결재 보조용 통계자료이며 원자료 수정용 파일이 아닙니다.",
    ]
    for note in notes:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
        c = ws.cell(row=row, column=1, value=note)
        c.font = Font(name="맑은 고딕", size=8, color="4B5563")
        c.alignment = Alignment(vertical="center", wrap_text=True)
        row += 1

    widths = [9, 7, 12, 15, 15, 11, 12, 15, 15, 12]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A9"
    ws.print_area = f"A1:J{row}"
    _setup_page(ws, landscape=True, repeat_rows=None)


def _build_after_school_sheet(wb: Workbook, state: dict[str, Any], as_of: str) -> None:
    ws = wb.create_sheet("방과후학습코칭")
    pivot = _after_school_pivot(state)
    ids = [x[0] for x in AFTER_SCHOOL_AREAS]
    labels = [x[1] for x in AFTER_SCHOOL_AREAS]
    max_col = 3 + len(ids)
    _merge_title(ws, "방과후학습코칭 지원현황", f"기준일: {_today_text(as_of)}", max_col)
    row = 5
    headers = ["학교급", "학년", *labels, "합계"]
    for col, value in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=value)
    _style_header_row(ws, row, 1, max_col)
    row += 1
    grand = {area: 0 for area in ids}
    for level, max_grade in (("초", 6), ("중", 3)):
        subtotal = {area: 0 for area in ids}
        for grade in range(1, max_grade + 1):
            data = pivot.get(f"{level}{grade}") or {}
            values = [int(data.get(area, 0) or 0) for area in ids]
            ws.append([level if grade == 1 else "", f"{grade}학년", *values, sum(values)])
            for area in ids:
                subtotal[area] += int(data.get(area, 0) or 0)
                grand[area] += int(data.get(area, 0) or 0)
        values = [subtotal[a] for a in ids]
        ws.append([f"{level} 소계", "", *values, sum(values)])
        _style_total_row(ws, ws.max_row, 1, max_col)
    values = [grand[a] for a in ids]
    ws.append(["합계", "", *values, sum(values)])
    _style_total_row(ws, ws.max_row, 1, max_col)
    _style_data_range(ws, 6, ws.max_row - 1, 1, max_col)
    for idx, width in enumerate([9, 9] + [13] * len(ids) + [10], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "C6"
    ws.print_area = f"A1:{get_column_letter(max_col)}{ws.max_row}"
    _setup_page(ws, landscape=True, repeat_rows="5:5")


def _build_therapy_sheet(wb: Workbook, state: dict[str, Any], as_of: str) -> None:
    ws = wb.create_sheet("치료기관연계")
    students, pivot = _therapy_data(state)
    _merge_title(ws, "치료기관연계 지원 현황", f"기준일: {_today_text(as_of)}  |  개인정보 최소화 명단", 8)
    row = 5
    headers = ["번호", "학교명", "학년", "성별", "지원내용", "연계기관", "지원기간", "비고"]
    for col, value in enumerate(headers, 1):
        ws.cell(row=row, column=col, value=value)
    _style_header_row(ws, row, 1, 8)
    row += 1
    for idx, student in enumerate(students, 1):
        therapy = student.get("therapy") or {}
        areas = student.get("areas") or []
        support = ", ".join(AREA_BY_ID.get(a, {"alt": a})["alt"] for a in areas)
        note = str(student.get("etcDetail") or "") if "ETC" in areas else ""
        ws.append([
            idx,
            student.get("sc") or "",
            f"{student.get('scType', '')}{student.get('gr', '')}",
            student.get("gen") or "",
            support,
            therapy.get("inst") or "",
            f"{therapy.get('start') or ''} ~ {therapy.get('end') or ''}",
            note,
        ])
    if not students:
        ws.append(["", "치료기관연계 대상 학생 없음", "", "", "", "", "", ""])
    _style_data_range(ws, 6, ws.max_row, 1, 8)

    row = ws.max_row + 3
    row = _section(ws, row, "지원내용별 집계", 8)
    headers2 = ["학교급", "학년", "난독증", "언어", "경계선지능", "ADHD", "심리정서", "기타"]
    for col, value in enumerate(headers2, 1):
        ws.cell(row=row, column=col, value=value)
    _style_header_row(ws, row, 1, 8)
    row += 1
    grand = {a: 0 for a in THERAPY_AREA_IDS}
    for level, max_grade in (("초", 6), ("중", 3)):
        subtotal = {a: 0 for a in THERAPY_AREA_IDS}
        for grade in range(1, max_grade + 1):
            data = pivot.get(f"{level}{grade}") or {}
            values = [int(data.get(a, 0) or 0) for a in THERAPY_AREA_IDS]
            ws.append([level if grade == 1 else "", f"{grade}학년", *values])
            for a in THERAPY_AREA_IDS:
                subtotal[a] += int(data.get(a, 0) or 0)
                grand[a] += int(data.get(a, 0) or 0)
        ws.append([f"{level} 소계", "", *[subtotal[a] for a in THERAPY_AREA_IDS]])
        _style_total_row(ws, ws.max_row, 1, 8)
    ws.append(["합계", "", *[grand[a] for a in THERAPY_AREA_IDS]])
    _style_total_row(ws, ws.max_row, 1, 8)
    _style_data_range(ws, row, ws.max_row - 1, 1, 8)
    for idx, width in enumerate([7, 17, 9, 8, 24, 20, 22, 20], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.freeze_panes = "A6"
    ws.print_area = f"A1:H{ws.max_row}"
    _setup_page(ws, landscape=True, repeat_rows="5:5")


def _build_dyslex_sheet(wb: Workbook, state: dict[str, Any], as_of: str) -> None:
    ws = wb.create_sheet("난독경계선")
    data = _dyslex_border(state)
    _merge_title(ws, "난독증 및 경계선지능 지원 현황", f"기준일: {_today_text(as_of)}", 10)
    row = 5
    top = ["분류", "지원신청", "진단검사", "난독증/경계선인 경우", "", "", "난독증/경계선 아닌 경우", "", "", "미지원 수 및 사유"]
    sub = ["", "", "", "치료지원", "학습코칭", "수업협력코칭", "치료지원", "학습코칭", "수업협력코칭", ""]
    for col, value in enumerate(top, 1):
        ws.cell(row=row, column=col, value=value)
    for col, value in enumerate(sub, 1):
        ws.cell(row=row + 1, column=col, value=value)
    ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=6)
    ws.merge_cells(start_row=row, start_column=7, end_row=row, end_column=9)
    for col in (1, 2, 3, 10):
        ws.merge_cells(start_row=row, start_column=col, end_row=row + 1, end_column=col)
    _style_header_row(ws, row, 1, 10)
    _style_header_row(ws, row + 1, 1, 10)
    row += 2
    for label, key in (("난독증", "dyslex"), ("경계선지능", "border")):
        d = data[key]
        reasons = ", ".join(dict.fromkeys(d["unsupReasons"])) or "-"
        ws.append([
            label, d["applied"], d["tested"], d["pos_therapy"], d["pos_coach"], d["pos_class"],
            d["neg_therapy"], d["neg_coach"], d["neg_class"], f"{d['unsup']} ({reasons})",
        ])
    _style_data_range(ws, row, ws.max_row, 1, 10)
    ws.merge_cells(start_row=ws.max_row + 2, start_column=1, end_row=ws.max_row + 2, end_column=10)
    note_cell = ws.cell(row=ws.max_row, column=1, value="※ 학생명은 포함하지 않으며, 통계 산출 기준은 프로그램의 등록·진단·지원유형 정보와 동일합니다.")
    note_cell.font = Font(name="맑은 고딕", size=8, color="4B5563")
    for idx, width in enumerate([12, 10, 10, 11, 11, 13, 11, 11, 13, 24], 1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.print_area = f"A1:J{ws.max_row}"
    _setup_page(ws, landscape=True, repeat_rows="5:6")


def create_statistics_report_xlsx(
    output_path: str | Path,
    state: dict[str, Any],
    *,
    as_of: str = "",
    report_type: str = "base",
) -> Path:
    statistics = build_statistics({"state": state})
    wb = Workbook()
    _build_summary_sheet(wb, state, statistics, as_of, report_type)
    _build_after_school_sheet(wb, state, as_of)
    _build_therapy_sheet(wb, state, as_of)
    _build_dyslex_sheet(wb, state, as_of)

    for ws in wb.worksheets:
        ws.sheet_properties.tabColor = NAVY if ws.title == "종합보고서" else None

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
