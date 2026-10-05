from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from services.document_routing import (
    execution_report_export_model as execution_report_model,
    manager_book_export_model as manager_book_model,
    pay_slip_export_model as pay_slip_model,
)


def _set_widths(ws, widths: list[int | float]) -> None:
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def create_pay_slip_xlsx(
    output_path: str | Path,
    settlement: dict[str, Any],
    staff: dict[str, Any],
    staff_id: str,
    ym: str,
    org: str = "",
) -> Path:
    model = pay_slip_model(settlement, staff, staff_id, ym, org)
    contract = model["contract"]
    wb = Workbook()
    ws = wb.active
    ws.title = contract["sheet"]
    ws.merge_cells("A1:H1")
    ws["A1"] = model["title"]
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    meta = dict(model["meta"])
    ws.append([f"수령인: {meta['수령인']}", "", f"소속: {meta['소속']}", "", f"발행일: {meta['발행일']}"])
    ws.append([])
    ws.append(model["columns"])
    for row in model["rows"]:
        ws.append(row)
    ws.append([])
    summary = model["summary"]
    ws.append([summary[0][0], summary[0][1], summary[1][0], summary[1][1], summary[2][0], summary[2][1]])
    ws.append([summary[3][0], summary[3][1], summary[4][0], summary[4][1], summary[5][0], summary[5][1]])
    _set_widths(ws, contract["columnWidths"])
    for row in ws.iter_rows(min_row=4):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def create_execution_xlsx(
    output_path: str | Path,
    settlement: dict[str, Any],
    staff_by_id: dict[str, dict[str, Any]],
    ym: str,
    org: str = "",
) -> Path:
    model = execution_report_model(settlement, staff_by_id, ym, org)
    contract = model["contract"]
    wb = Workbook()
    ws = wb.active
    ws.title = contract["sheet"]
    ws.merge_cells("A1:J1")
    ws["A1"] = model["title"]
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    meta = dict(model["meta"])
    ws.append([f"기관: {meta['기관']}", "", f"대상: {meta['대상']}", "", f"발행일: {meta['발행일']}"])
    ws.append([])
    ws.append(model["columns"])
    for row in model["rows"]:
        ws.append(row)
    ws.append([])
    ws.append(model["summaryRow"])
    _set_widths(ws, contract["columnWidths"])
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


def create_manager_book_xlsx(output_path: str | Path, state: dict[str, Any], staff_id: str, ym: str) -> Path:
    model = manager_book_model(state, staff_id, ym)
    contract = model["contract"]
    wb = Workbook()
    ws = wb.active
    ws.title = contract["sheet"]
    ws.merge_cells("A1:M1")
    ws["A1"] = model["title"]
    ws["A1"].font = Font(bold=True, size=16)
    ws["A1"].alignment = Alignment(horizontal="center")
    meta = dict(model["meta"])
    ws.append([f"소속: {meta['소속']}", "", f"성명: {meta['성명']}", "", f"연락처: {meta['연락처']}"])
    ws.append([f"담당학생: {meta['담당학생']}"])
    ws.append([])
    ws.append(model["columns"])
    for row in model["rows"]:
        ws.append(row)
    ws.append([])
    totals = model["totals"]
    ws.append(["", "", "", "", f"총 실시 회기: {totals['count']}회 · 총 시수: {totals['hours']:.1f}시간 · 확인자({totals['confirmer']}):", "(인)"])
    _set_widths(ws, contract["columnWidths"])
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
